"""End to end: generate data with planted problems, run the pipeline, and check
that it catches what was planted without raising false alarms."""
import tempfile
import unittest
from datetime import date
from pathlib import Path

from openpyxl import load_workbook

from dim.generate import generate
from dim.pipeline import run


class EndToEnd(unittest.TestCase):
    def test_detection_against_answer_key(self):
        for seed in (1, 2, 3):
            with self.subTest(seed=seed), tempfile.TemporaryDirectory() as tmp:
                data, out = Path(tmp, "data"), Path(tmp, "out")
                generate(data, n=2000, seed=seed)
                m = run(data, out, date(2026, 10, 1))
                for row in m["evaluation"]:
                    if row["check"] == "DUP":
                        self.assertGreaterEqual(row["recall"], 0.9, row)
                        self.assertGreaterEqual(row["precision"], 0.9, row)
                    elif row["injected"]:
                        self.assertEqual(row["recall"], 1.0, row)
                        self.assertEqual(row["precision"], 1.0, row)

    def test_outputs_written(self):
        with tempfile.TemporaryDirectory() as tmp:
            data, out = Path(tmp, "data"), Path(tmp, "out")
            generate(data, n=500, seed=5)
            m = run(data, out, date(2026, 10, 1))
            for name in ["dashboard.html", "exceptions.csv", "duplicate_review.csv", "metrics.json",
                         "run_history.csv", "exception_report_2026-10-01.xlsx"]:
                self.assertTrue((out / name).exists(), name)
            wb = load_workbook(out / "exception_report_2026-10-01.xlsx")
            self.assertEqual(wb.sheetnames, ["Summary", "Admissions", "Registrar", "Integration",
                                             "Duplicate Review", "Detection Check"])
            self.assertGreater(m["quality_score"], 0.8)
            self.assertLess(m["sync_rate"], 1.0)

    def test_history_appends_one_row_per_run_date(self):
        with tempfile.TemporaryDirectory() as tmp:
            data, out = Path(tmp, "data"), Path(tmp, "out")
            generate(data, n=300, seed=9)
            for d in [date(2026, 10, 1), date(2026, 10, 2), date(2026, 10, 2)]:
                run(data, out, d, write_reports=False)
            lines = (out / "run_history.csv").read_text().strip().splitlines()
            self.assertEqual(len(lines), 3)  # header + 2 distinct dates


if __name__ == "__main__":
    unittest.main()
