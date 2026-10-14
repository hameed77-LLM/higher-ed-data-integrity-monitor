"""Outputs: an Excel exception report split by owning office, and a self-contained HTML dashboard."""
from __future__ import annotations

import html
import json
from pathlib import Path

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HEADER_FILL = PatternFill("solid", fgColor="1F3A5F")
HEADER_FONT = Font(bold=True, color="FFFFFF")
SEVERITY_FILL = {"High": "F8D7D7", "Medium": "FDEBC8", "Low": "E4EEF9"}


def _format_sheet(ws, widths: dict[str, int] | None = None) -> None:
    for cell in ws[1]:
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
        cell.alignment = Alignment(vertical="center")
    ws.freeze_panes = "A2"
    if ws.max_row > 1 and ws.title != "Summary":
        ws.auto_filter.ref = ws.dimensions
    for idx, col in enumerate(ws.iter_cols(min_row=1, max_row=min(ws.max_row, 200)), start=1):
        letter = get_column_letter(idx)
        longest = max(len(str(c.value)) if c.value is not None else 0 for c in col)
        ws.column_dimensions[letter].width = (widths or {}).get(col[0].value, min(max(10, longest + 2), 60))


def _shade_severity(ws) -> None:
    headers = [c.value for c in ws[1]]
    if "severity" not in headers:
        return
    col = headers.index("severity") + 1
    for row in ws.iter_rows(min_row=2):
        fill = SEVERITY_FILL.get(row[col - 1].value)
        if fill:
            row[col - 1].fill = PatternFill("solid", fgColor=fill)


def write_excel(path: Path, metrics: dict, by_rule: pd.DataFrame, exceptions: pd.DataFrame,
                duplicates: list[dict], evaluation: list[dict]) -> None:
    summary = pd.DataFrame([
        ("Run date", metrics["run_date"]),
        ("Records checked", metrics["records_checked"]),
        ("  Slate applicants", metrics["slate_applicants"]),
        ("  Colleague students", metrics["colleague_records"]),
        ("Rules run", metrics["rules"]),
        ("Open exceptions", metrics["open_exceptions"]),
        ("  High severity", metrics["high_severity"]),
        ("Records with at least one issue", metrics["records_flagged"]),
        ("Data quality score (clean records)", f"{metrics['quality_score']:.1%}"),
        ("Possible duplicate applicants", metrics["possible_duplicates"]),
        ("Deposited students synced to Colleague", f"{metrics['deposited_in_colleague']} of {metrics['deposited']}"),
        ("Birth dates normalized on load", metrics["dates_normalized"]),
    ], columns=["Measure", "Value"])
    cols = ["rule_id", "rule", "severity", "record_id", "detail"]
    with pd.ExcelWriter(path, engine="openpyxl") as xw:
        summary.to_excel(xw, sheet_name="Summary", index=False)
        by_rule.to_excel(xw, sheet_name="Summary", index=False, startrow=len(summary) + 2)
        for owner in ["Admissions", "Registrar", "Integration"]:
            part = exceptions[exceptions.owner == owner][cols].sort_values(["rule_id", "record_id"])
            part.to_excel(xw, sheet_name=owner, index=False)
        pd.DataFrame(duplicates, columns=["record_a", "name_a", "record_b", "name_b", "score", "band", "evidence"]) \
            .to_excel(xw, sheet_name="Duplicate Review", index=False)
        if evaluation:
            pd.DataFrame(evaluation).to_excel(xw, sheet_name="Detection Check", index=False)
        for ws in xw.book.worksheets:
            _format_sheet(ws, {"detail": 55, "evidence": 60, "rule": 42, "Measure": 40})
            _shade_severity(ws)
        ws = xw.book["Summary"]
        hdr = len(summary) + 3
        for cell in ws[hdr]:
            cell.fill, cell.font = HEADER_FILL, HEADER_FONT


# ----------------------------------------------------------------------------- dashboard

def _line_svg(history: pd.DataFrame, col: str, fmt) -> str:
    if len(history) < 2:
        return '<p class="muted">The trend appears after two or more runs.</p>'
    w, h, left, right, top, bottom = 640, 220, 48, 16, 14, 30
    vals = history[col].astype(float).tolist()
    dates = history.run_date.tolist()
    lo, hi = 0.0, max(vals) * 1.1 or 1
    x = lambda i: left + i * (w - left - right) / (len(vals) - 1)
    y = lambda v: top + (1 - (v - lo) / (hi - lo)) * (h - top - bottom)
    grid = "".join(
        f'<line x1="{left}" x2="{w - right}" y1="{y(hi * k / 4):.1f}" y2="{y(hi * k / 4):.1f}" class="grid"/>'
        f'<text x="{left - 6}" y="{y(hi * k / 4) + 4:.1f}" class="tick" text-anchor="end">{fmt(hi * k / 4)}</text>'
        for k in range(5))
    pts = " ".join(f"{x(i):.1f},{y(v):.1f}" for i, v in enumerate(vals))
    hits = "".join(
        f'<circle cx="{x(i):.1f}" cy="{y(v):.1f}" r="10" class="hit" data-tip="{d}: {fmt(v)}"/>'
        f'<circle cx="{x(i):.1f}" cy="{y(v):.1f}" r="4" class="dot"/>'
        for i, (d, v) in enumerate(zip(dates, vals)))
    labels = "".join(f'<text x="{x(i):.1f}" y="{h - 8}" class="tick" text-anchor="middle">{dates[i][5:]}</text>'
                     for i in sorted({0, len(vals) // 2, len(vals) - 1}))
    return (f'<svg viewBox="0 0 {w} {h}" role="img" aria-label="trend">{grid}'
            f'<polyline points="{pts}" class="line"/>{hits}{labels}</svg>')


def _bar_svg(by_rule: pd.DataFrame) -> str:
    rows = by_rule.sort_values("count", ascending=False)
    bar_h, gap, left, right = 18, 8, 300, 40
    w = 1000
    h = len(rows) * (bar_h + gap)
    top = max(rows["count"].max(), 1)
    out = []
    for i, r in enumerate(rows.itertuples()):
        yy = i * (bar_h + gap)
        bw = max(2, (w - left - right) * r.count / top) if r.count else 0
        label = html.escape(f"{r.rule_id}  {r.rule}")
        out.append(
            f'<text x="{left - 8}" y="{yy + 13}" class="lbl" text-anchor="end">{label}</text>'
            f'<rect x="{left}" y="{yy}" width="{bw:.1f}" height="{bar_h}" rx="4" class="bar"/>'
            f'<rect x="0" y="{yy - gap / 2}" width="{w}" height="{bar_h + gap}" class="hit" '
            f'data-tip="{label}: {r.count} ({r.owner}, {r.severity})"/>'
            f'<text x="{left + bw + 6:.1f}" y="{yy + 13}" class="val">{r.count}</text>')
    return f'<svg viewBox="0 0 {w} {h}" role="img" aria-label="exceptions by rule">{"".join(out)}</svg>'


def write_dashboard(path: Path, metrics: dict, by_rule: pd.DataFrame, history: pd.DataFrame,
                    duplicates: list[dict], evaluation: list[dict]) -> None:
    pct = lambda v: f"{v:.1%}"
    tiles = [
        ("Records checked", f"{metrics['records_checked']:,}",
         f"{metrics['slate_applicants']:,} Slate + {metrics['colleague_records']:,} Colleague"),
        ("Open exceptions", f"{metrics['open_exceptions']:,}", f"{metrics['high_severity']} high severity"),
        ("Data quality score", pct(metrics["quality_score"]), "records with no open issue"),
        ("Possible duplicates", f"{metrics['possible_duplicates']}", f"{metrics['definite_duplicates']} definite"),
        ("Slate → Colleague sync", pct(metrics["sync_rate"]),
         f"{metrics['deposited_in_colleague']} of {metrics['deposited']} deposits"),
    ]
    tiles_html = "".join(f'<div class="tile"><div class="k">{k}</div><div class="v">{v}</div>'
                         f'<div class="s">{s}</div></div>' for k, v, s in tiles)
    owners = metrics["by_owner"]
    owner_html = "".join(f'<div class="owner"><span>{o}</span><b>{owners.get(o, 0):,}</b></div>'
                         for o in ["Admissions", "Registrar", "Integration"])
    eval_map = {e["check"]: e for e in evaluation}
    rule_rows = "".join(
        f'<tr><td>{r.rule_id}</td><td>{html.escape(r.rule)}</td><td>{r.owner}</td>'
        f'<td><span class="sev {r.severity.lower()}">{r.severity}</span></td><td class="num">{r.count}</td>'
        + (f'<td class="num">{eval_map[r.rule_id]["caught"]}/{eval_map[r.rule_id]["injected"]}</td>' if evaluation else "")
        + "</tr>"
        for r in by_rule.itertuples())
    dup_rows = "".join(
        f'<tr><td>{d["record_a"]}</td><td>{html.escape(d["name_a"])}</td><td>{d["record_b"]}</td>'
        f'<td>{html.escape(d["name_b"])}</td><td class="num">{d["score"]:.2f}</td>'
        f'<td><span class="band {d["band"]}">{d["band"]}</span></td><td>{html.escape(d["evidence"])}</td></tr>'
        for d in sorted(duplicates, key=lambda d: (d["band"] != "likely", -d["score"]))[:40])
    dup_eval = eval_map.get("DUP")
    eval_note = ""
    if evaluation:
        caught = sum(e["caught"] for e in evaluation)
        injected = sum(e["injected"] for e in evaluation)
        flagged = sum(e["flagged"] for e in evaluation)
        eval_note = (f'<p class="muted">Detection check against the generator\'s answer key: caught '
                     f'{caught} of {injected} injected issues with {flagged - caught} false alarms'
                     + (f'; duplicate matcher found {dup_eval["caught"]} of {dup_eval["injected"]} planted pairs.' if dup_eval else ".")
                     + '</p>')

    page = f"""<!doctype html>
<html lang="en"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1">
<title>Data Integrity Monitor</title>
<style>
:root {{ --surface:#fcfcfb; --panel:#ffffff; --border:#e4e2dc; --text:#0b0b0b; --text2:#52514e; --muted:#7a7972;
  --series:#2a78d6; --grid:#ecebe7; --high:#d03b3b; --med:#b07800; --low:#2a78d6; color-scheme:light; }}
@media (prefers-color-scheme: dark) {{ :root:not([data-theme="light"]) {{ --surface:#1a1a19; --panel:#222220;
  --border:#383835; --text:#f0efec; --text2:#c3c2b7; --muted:#97968c; --series:#3987e5; --grid:#2e2e2b;
  --high:#e66767; --med:#e0a630; --low:#6da7ec; color-scheme:dark; }} }}
:root[data-theme="dark"] {{ --surface:#1a1a19; --panel:#222220; --border:#383835; --text:#f0efec; --text2:#c3c2b7;
  --muted:#97968c; --series:#3987e5; --grid:#2e2e2b; --high:#e66767; --med:#e0a630; --low:#6da7ec; color-scheme:dark; }}
* {{ box-sizing:border-box; }}
body {{ margin:0; background:var(--surface); color:var(--text); font:14px/1.45 system-ui,-apple-system,"Segoe UI",sans-serif; }}
main {{ max-width:1080px; margin:0 auto; padding:24px 16px 48px; }}
h1 {{ font-size:22px; margin:0 0 2px; }} h2 {{ font-size:15px; margin:0 0 12px; }}
.sub, .muted {{ color:var(--muted); }} .sub {{ margin:0 0 20px; }}
.tiles {{ display:grid; grid-template-columns:repeat(auto-fit,minmax(170px,1fr)); gap:12px; margin-bottom:16px; }}
.tile, .card {{ background:var(--panel); border:1px solid var(--border); border-radius:10px; padding:14px 16px; }}
.tile .k {{ color:var(--text2); font-size:12px; }} .tile .v {{ font-size:26px; font-weight:650; margin:2px 0; }}
.tile .s {{ color:var(--muted); font-size:12px; }}
.grid2 {{ display:grid; grid-template-columns:1fr; gap:16px; margin-bottom:16px; }}
@media (min-width:900px) {{ .grid2 {{ grid-template-columns:3fr 2fr; }} }}
.owners {{ display:flex; flex-direction:column; gap:8px; }}
.owner {{ display:flex; justify-content:space-between; border-bottom:1px solid var(--border); padding:6px 0; }}
.owner b {{ font-variant-numeric:tabular-nums; }}
svg {{ width:100%; height:auto; display:block; overflow:visible; }}
.grid {{ stroke:var(--grid); stroke-width:1; }} .tick, .lbl, .val {{ fill:var(--text2); font-size:11px; }}
.val {{ fill:var(--text); font-variant-numeric:tabular-nums; }}
.line {{ fill:none; stroke:var(--series); stroke-width:2; }} .dot {{ fill:var(--series); stroke:var(--panel); stroke-width:2; pointer-events:none; }}
.bar {{ fill:var(--series); }} .hit {{ fill:transparent; cursor:default; }} .hit:hover + .dot {{ r:6; }}
rect.hit:hover {{ fill:var(--grid); fill-opacity:.5; }}
.scroll {{ overflow-x:auto; }}
table {{ border-collapse:collapse; width:100%; font-size:13px; }}
th, td {{ text-align:left; padding:7px 8px; border-bottom:1px solid var(--border); vertical-align:top; }}
th {{ color:var(--text2); font-weight:600; font-size:12px; }} td.num, th.num {{ text-align:right; font-variant-numeric:tabular-nums; }}
.sev, .band {{ font-size:11px; padding:1px 7px; border-radius:99px; border:1px solid currentColor; white-space:nowrap; }}
.sev.high {{ color:var(--high); }} .sev.medium {{ color:var(--med); }} .sev.low {{ color:var(--low); }}
.band.definite {{ color:var(--high); }} .band.likely {{ color:var(--med); }}
.card {{ margin-bottom:16px; }}
#tip {{ position:fixed; pointer-events:none; background:var(--text); color:var(--surface); font-size:12px;
  padding:5px 8px; border-radius:6px; opacity:0; transition:opacity .1s; max-width:280px; }}
</style></head>
<body><main>
<h1>Data Integrity Monitor</h1>
<p class="sub">Nightly run {metrics['run_date']} · {metrics['rules']} rules across Slate (admissions) and Colleague (student records) · synthetic demo data</p>
<div class="tiles">{tiles_html}</div>
<div class="grid2">
  <div class="card"><h2>Open exceptions per nightly run</h2>{_line_svg(history, 'open_exceptions', lambda v: f'{v:,.0f}')}</div>
  <div class="card"><h2>Exceptions by owning office</h2><div class="owners">{owner_html}</div>
    <p class="muted" style="margin-top:12px">Each office gets its own tab in the Excel exception report.</p></div>
</div>
<div class="card"><h2>Exceptions by rule</h2><div class="scroll"><div style="min-width:640px">{_bar_svg(by_rule)}</div></div></div>
<div class="card"><h2>Rule catalog</h2>{eval_note}<div class="scroll"><table>
<tr><th>Rule</th><th>Check</th><th>Owner</th><th>Severity</th><th class="num">Open</th>{'<th class="num">Caught / planted</th>' if evaluation else ''}</tr>
{rule_rows}</table></div></div>
<div class="card"><h2>Duplicate review queue (top {min(40, len(duplicates))} of {len(duplicates)}, likely matches first)</h2><div class="scroll"><table>
<tr><th>Record A</th><th>Name</th><th>Record B</th><th>Name</th><th class="num">Score</th><th>Band</th><th>Evidence</th></tr>
{dup_rows}</table></div></div>
</main><div id="tip"></div>
<script>
const tip = document.getElementById('tip');
document.querySelectorAll('[data-tip]').forEach(el => {{
  el.addEventListener('mousemove', e => {{ tip.textContent = el.dataset.tip; tip.style.opacity = 1;
    tip.style.left = Math.min(e.clientX + 12, innerWidth - 290) + 'px'; tip.style.top = (e.clientY + 12) + 'px'; }});
  el.addEventListener('mouseleave', () => tip.style.opacity = 0);
}});
</script>
<script type="application/json" id="metrics">{html.escape(json.dumps(metrics))}</script>
</body></html>"""
    path.write_text(page, encoding="utf-8")
